from PyQt6 import QtWidgets #from PyQt5 import QtWidgets
import design
import sys
import os
import csv
import json
import PyPDF2
import docx
import openpyxl


class MainWindow(QtWidgets.QMainWindow, design.Ui_MainWindow):
    def __init__(self):
        super().__init__()
        self.setupUi(self)
        self.pushButton.clicked.connect(self.change_directory)
        self.pushButton_2.clicked.connect(self.start)
        self.label.setOpenExternalLinks(True)
        self.listWidget.addItem(
            'Hello!\n1. Choose a directory\n2. Select file extension\n3. Enter the search word\n4. Press start')
        self.main_directory = None
        self.extension_list = []
        self.dir_list = []
        self.encoding_list = ['utf-8', 'cp1251', 'utf-16']
        self.result = []
        self.keyword = None

    def change_directory(self):
        self.main_directory = QtWidgets.QFileDialog.getExistingDirectory(self, 'Change directory')
        self.textEdit.setPlainText(self.main_directory)

    def start(self):
        self.keyword = self.textEdit_2.toPlainText()
        if len(self.keyword) == 0:
            return
        self.listWidget.clear()
        self.extension_list = []
        self.dir_list = []
        self.result = []
        if self.main_directory:
            self.check_extension()
            if self.extension_list:
                self.listWidget.addItem('Step 1/3:\nFile search')
                self.listWidget.repaint()
                for root, dirs, files in os.walk(self.main_directory):
                    for file in files:
                        for extension in self.extension_list:
                            if file.endswith(extension):
                                file_full_name = os.path.join(root, file)
                                self.dir_list.append(file_full_name)
                if self.dir_list:
                    self.listWidget.addItem(f'Step 2/3:\n{len(self.dir_list)} files found. Cheking found files')
                    self.listWidget.repaint()
                    self.search_file()
                else:
                    self.listWidget.addItem('Files with selected extensions not found')
            else:
                self.textEdit.setPlainText('Select file extension, please')
        else:
            self.textEdit.setPlainText('Change directory, please')

    def check_extension(self):
        if self.checkBox_txt.isChecked():
            self.extension_list.append('.txt')
        if self.checkBox_docx.isChecked():
            self.extension_list.append('.docx')
        if self.checkBox_xlsx.isChecked():
            self.extension_list.append('.xlsx')
        if self.checkBox_json.isChecked():
            self.extension_list.append('.json')
        if self.checkBox_csv.isChecked():
            self.extension_list.append('.csv')
        if self.checkBox_pdf.isChecked():
            self.extension_list.append('.pdf')

    def search_file(self):
        for file_full_name in self.dir_list:
            if file_full_name in self.result:
                continue
            try:
                if file_full_name.endswith('.txt'):
                    self.read_txt(file_full_name)
                elif file_full_name.endswith('.pdf'):
                    self.read_pdf(file_full_name)
                elif file_full_name.endswith('.xlsx'):
                    self.read_xlsx(file_full_name)
                elif file_full_name.endswith('.docx'):
                    self.read_docx(file_full_name)
                elif file_full_name.endswith('.json'):
                    self.read_json(file_full_name)
                elif file_full_name.endswith('.csv'):
                    self.read_csv(file_full_name)
            except:
                pass
        if self.result:
            self.listWidget.addItem(f'Step 3/3:\nMatches in {len(self.result)} files:')
            self.listWidget.repaint()
            for file in self.result:
                self.listWidget.addItem(file)

        else:
            self.listWidget.addItem('Step 3/3:\nResults\nNot found :(')
            self.listWidget.repaint()

    def read_txt(self, file_full_name):
        for encoding in self.encoding_list:
            with open(file_full_name, 'r', encoding=encoding) as file:
                for line in file:
                    if self.keyword.lower() in line.lower():
                        return self.result.append(file_full_name)

    def read_pdf(self, file_full_name):
        with open(file_full_name, 'rb') as pdf_file:
            pdf_reader = PyPDF2.PdfFileReader(pdf_file)
            for page_num in range(pdf_reader.numPages):
                pdf_page = pdf_reader.getPage(page_num)
                if self.keyword.lower() in pdf_page.extractText().lower():
                    return self.result.append(file_full_name)

    def read_xlsx(self, file_full_name):
        book = openpyxl.open(file_full_name, read_only=True)
        sheet = book.active
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value != None:
                    if self.keyword in str(cell.value):
                        return self.result.append(file_full_name)

    def read_docx(self, file_full_name):
        doc = docx.Document(file_full_name)

    def read_json(self, file_full_name):
        for paragraph in doc.paragraphs:
            if self.keyword.lower() in paragraph.text.lower():
                return self.result.append(file_full_name)
        with open(file_full_name) as file:
            file_reader = json.load(file)
            for keys, values in file_reader.items():
                if self.keyword in keys:
                    return self.result.append(file_full_name)
                elif self.keyword in values:
                    return self.result.append(file_full_name)

    def read_csv(self, file_full_name):
        with open(file_full_name, 'r') as csv_file:
            file_reader = csv.reader(csv_file)
            for row in file_reader:
                if self.keyword in row:
                    return self.result.append(file_full_name)


def main():
    app = QtWidgets.QApplication(sys.argv)
    window = MainWindow()
    window.show()
    app.exec()


if __name__ == '__main__':
    main()
